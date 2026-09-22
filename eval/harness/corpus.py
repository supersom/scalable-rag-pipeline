import argparse
import hashlib
import json
import random
import shutil
import sys
from pathlib import Path

DEFAULT_HARD_KEYWORDS = ("scheduling", "distributed", "container", "autoscal", "cluster", "kubernetes",
                         "orchestrat", "docker", "cloud", "load balanc")
POOL_DIRS = ("eval/datasets/noisy_data", "eval/datasets/ingested", "eval/datasets/digested")
TRUE_DIR = "eval/datasets/true_data"
STAGE_DIR = "eval/datasets/eval_corpus"
MANIFEST_PATH = "eval/sets/corpus_manifest.json"
_GROUPS = ("forced", "hard_negatives", "random")


def sha256_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def list_pool(dirs, suffixes=(".pdf",)):
    files = [p for d in dirs if Path(d).is_dir() for p in Path(d).iterdir()
             if p.is_file() and p.suffix.lower() in suffixes]
    files.sort(key=lambda p: p.name)
    names = [p.name for p in files]
    duplicates = sorted({n for n in names if names.count(n) > 1})
    if duplicates:
        raise ValueError(f"duplicate filenames in the pool: {duplicates}")
    return files


def select_sample(pool, *, seed, n_random, n_hard, keywords=DEFAULT_HARD_KEYWORDS, include=()):
    names = [p.name for p in pool]
    unknown = [n for n in include if n not in names]
    if unknown:
        raise ValueError(f"--include names not in the pool: {unknown}")
    rng = random.Random(seed)
    forced = sorted(include)
    remaining = [n for n in names if n not in forced]
    candidates = [n for n in remaining if any(k in n.lower() for k in keywords)]
    hard = sorted(rng.sample(candidates, min(n_hard, len(candidates))))
    rest = [n for n in remaining if n not in hard]
    return {"forced": forced, "hard_negatives": hard, "random": sorted(rng.sample(rest, min(n_random, len(rest))))}


def build_corpus_manifest(true_dir, pool, selection, seed):
    by_name = {p.name: p for p in pool}
    manifest = {"seed": seed, "true_data": [{"filename": p.name, "sha256": sha256_file(p)}
                                            for p in sorted(Path(true_dir).iterdir()) if p.is_file()]}
    for group in _GROUPS:
        manifest[group] = [{"filename": n, "sha256": sha256_file(by_name[n])} for n in selection[group]]
    return manifest


def stage_corpus(manifest, true_dir, pool, dest):
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    by_name = {p.name: p for p in pool}
    for entry in manifest["true_data"]:
        _copy_checked(Path(true_dir) / entry["filename"], entry, dest)
    for group in _GROUPS:
        for entry in manifest[group]:
            _copy_checked(by_name[entry["filename"]], entry, dest)
    return sum(1 for p in dest.iterdir() if p.is_file())


def _copy_checked(src, entry, dest):
    if sha256_file(src) != entry["sha256"]:
        raise ValueError(f"{entry['filename']} changed since the corpus manifest was built; rebuild the manifest")
    shutil.copy2(src, dest / entry["filename"])


def extract_text(path, max_pages=3):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return " ".join((page.extract_text() or "") for page in PdfReader(str(path)).pages[:max_pages])
    if suffix == ".docx":
        from docx import Document

        return "\n".join(p.text for p in Document(str(path)).paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation

        return "\n".join(shape.text_frame.text for slide in Presentation(str(path)).slides
                         for shape in slide.shapes if shape.has_text_frame)
    if suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup

        return BeautifulSoup(path.read_text(errors="ignore"), "html.parser").get_text(" ")
    return path.read_text(errors="ignore")


def main(argv=None):
    parser = argparse.ArgumentParser(prog="python -m eval.harness.corpus")
    commands = parser.add_subparsers(dest="command", required=True)
    build = commands.add_parser("build")
    build.add_argument("--seed", type=int, default=42)
    build.add_argument("--n-random", type=int, default=60)
    build.add_argument("--n-hard", type=int, default=20)
    build.add_argument("--include", nargs="*", default=[])
    build.add_argument("--out", default=MANIFEST_PATH)
    stage = commands.add_parser("stage")
    stage.add_argument("--manifest", default=MANIFEST_PATH)
    stage.add_argument("--dest", default=STAGE_DIR)
    args = parser.parse_args(argv)

    pool = list_pool(POOL_DIRS)
    if args.command == "build":
        selection = select_sample(pool, seed=args.seed, n_random=args.n_random, n_hard=args.n_hard,
                                  include=args.include)
        manifest = build_corpus_manifest(TRUE_DIR, pool, selection, args.seed)
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")
        print(f"pool {len(pool)} pdfs; manifest {out}: " + ", ".join(f"{g} {len(manifest[g])}" for g in ("true_data", *_GROUPS)))
    else:
        manifest = json.loads(Path(args.manifest).read_text())
        print(f"staged {stage_corpus(manifest, TRUE_DIR, pool, args.dest)} files in {args.dest}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
